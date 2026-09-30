import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

def generate_manual_test_cases():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    # Define Header Columns
    headers = [
        "Test Case ID",
        "Module",
        "Test Scenario",
        "Preconditions",
        "Test Steps",
        "Test Data",
        "Expected Result",
        "Actual Result",
        "Status",
        "Severity",
        "Priority"
    ]

    # Test Cases Data (45 detailed test cases)
    test_cases = [
        # LOGIN MODULE
        ("TC_LOG_001", "Login", "Valid Admin Login", "User is on /login page", "1. Enter admin email\n2. Enter valid password\n3. Click Login button", "admin@workforce.test / Admin@12345", "Redirect to /dashboard with Admin badge visible", "Redirected to /dashboard successfully", "Pass", "High", "P1"),
        ("TC_LOG_002", "Login", "Valid HR Manager Login", "User is on /login page", "1. Enter HR email\n2. Enter valid password\n3. Click Login button", "hr@workforce.test / HR@12345", "Redirect to /dashboard with HR Manager badge visible", "Redirected to /dashboard successfully", "Pass", "High", "P1"),
        ("TC_LOG_003", "Login", "Valid Employee Login", "User is on /login page", "1. Enter employee email\n2. Enter valid password\n3. Click Login button", "employee@workforce.test / Employee@12345", "Redirect to /profile with Employee badge visible", "Redirected to /profile successfully", "Pass", "High", "P1"),
        ("TC_LOG_004", "Login", "Login with Invalid Password", "User is on /login page", "1. Enter valid email\n2. Enter incorrect password\n3. Click Login", "admin@workforce.test / WrongPass123", "Display error message 'Invalid email or password'", "Displayed 'Invalid email or password'", "Pass", "High", "P1"),
        ("TC_LOG_005", "Login", "Login with Non-existent Email", "User is on /login page", "1. Enter non-existent email\n2. Enter password\n3. Click Login", "unknown@workforce.test / Admin@12345", "Display error message 'Invalid email or password'", "Displayed 'Invalid email or password'", "Pass", "Medium", "P2"),
        ("TC_LOG_006", "Login", "Login with Empty Email", "User is on /login page", "1. Leave email empty\n2. Enter password\n3. Click Login", "Password: Admin@12345", "Display error message 'Email is required'", "Displayed 'Email is required'", "Pass", "Medium", "P2"),
        ("TC_LOG_007", "Login", "Login with Empty Password", "User is on /login page", "1. Enter email\n2. Leave password empty\n3. Click Login", "Email: admin@workforce.test", "Display error message 'Password is required'", "Displayed 'Password is required'", "Pass", "Medium", "P2"),
        ("TC_LOG_008", "Login", "Login with Invalid Email Format", "User is on /login page", "1. Enter invalid email format (no domain/@)\n2. Enter password\n3. Click Login", "admin_workforce.test / Admin@12345", "Display error message 'Invalid email format'", "Displayed 'Invalid email format'", "Pass", "Low", "P3"),
        ("TC_LOG_009", "Login", "User Logout Functionality", "User is logged in", "1. Click Logout button in navbar", "N/A", "Session ended, token cookie removed, redirect to /login", "Redirected to /login with logout message", "Pass", "High", "P1"),

        # DASHBOARD MODULE
        ("TC_DSH_010", "Dashboard", "Admin Dashboard Initial Render", "Admin logged in", "1. Navigate to /dashboard", "N/A", "Page loads with summary cards and recent employee table", "Summary cards and table rendered cleanly", "Pass", "High", "P1"),
        ("TC_DSH_011", "Dashboard", "Total Employees Metric Accuracy", "Admin logged in", "1. View Total Employees metric card", "N/A", "Count matches total employee records in DB (15+)", "Displayed exact DB total count", "Pass", "Medium", "P2"),
        ("TC_DSH_012", "Dashboard", "Active vs Inactive Metrics Breakdown", "Admin logged in", "1. Check Active and Inactive count cards", "N/A", "Active + Inactive equals Total Employees", "Active + Inactive equals total count", "Pass", "Medium", "P2"),
        ("TC_DSH_013", "Dashboard", "Recent Employees Table Display", "Admin logged in", "1. Inspect Recent Employees table", "N/A", "Displays top 5 most recently created employees", "Rendered top 5 recent records", "Pass", "Low", "P3"),
        ("TC_DSH_014", "Dashboard", "Quick Add Employee Button", "Admin logged in", "1. Click '+ Add Employee' button on dashboard", "N/A", "Redirect to /employees/add form page", "Redirected to /employees/add page", "Pass", "Medium", "P2"),

        # EMPLOYEE DIRECTORY & SEARCH/FILTER MODULE
        ("TC_EMP_015", "Employees", "View Employee Directory Table", "Admin/HR logged in", "1. Navigate to /employees", "N/A", "Table lists all employees with columns & action buttons", "Table displayed with correct columns", "Pass", "High", "P1"),
        ("TC_EMP_016", "Employees", "Search Employee by First Name", "On /employees page", "1. Type 'Alexander' in search input", "Query: Alexander", "Table filters to show only Alexander Wright", "Filtered correctly to 1 row", "Pass", "High", "P1"),
        ("TC_EMP_017", "Employees", "Search Employee by Email", "On /employees page", "1. Type 'sarah.jenkins@workforce.test'", "Query: sarah.jenkins@workforce.test", "Table filters to show Sarah Jenkins", "Filtered correctly", "Pass", "High", "P1"),
        ("TC_EMP_018", "Employees", "Search Employee by Code", "On /employees page", "1. Type 'EMP003' in search input", "Query: EMP003", "Table filters to show John Doe (EMP003)", "Filtered correctly", "Pass", "High", "P1"),
        ("TC_EMP_019", "Employees", "Search with Non-Matching Term", "On /employees page", "1. Type 'XYZ999' in search input", "Query: XYZ999", "Table rows hidden / empty state message", "No matching rows displayed", "Pass", "Medium", "P2"),
        ("TC_EMP_020", "Employees", "Filter Employees by Department", "On /employees page", "1. Select 'Engineering' from Department dropdown", "Dept: Engineering", "Table displays only Engineering department staff", "Filtered to Engineering staff", "Pass", "High", "P1"),
        ("TC_EMP_021", "Employees", "Filter Employees by Active Status", "On /employees page", "1. Select 'Active' from Status dropdown", "Status: Active", "Table displays only Active employees", "Filtered to Active staff", "Pass", "High", "P1"),
        ("TC_EMP_022", "Employees", "Filter Employees by Inactive Status", "On /employees page", "1. Select 'Inactive' from Status dropdown", "Status: Inactive", "Table displays only Inactive employees", "Filtered to Inactive staff", "Pass", "High", "P1"),
        ("TC_EMP_023", "Employees", "Filter Employees by Job Position", "On /employees page", "1. Type 'Architect' in position filter", "Position: Architect", "Table displays matching job positions", "Filtered to matching position", "Pass", "Medium", "P2"),
        ("TC_EMP_024", "Employees", "Combined Search and Department Filter", "On /employees page", "1. Select 'Engineering' dept\n2. Type 'John' in search", "Dept: Engineering, Query: John", "Table displays John Doe (Engineering)", "Filtered accurately", "Pass", "Medium", "P2"),

        # CRUD & VALIDATION MODULE
        ("TC_CRUD_025", "Add/Edit", "Add New Employee with Valid Data", "Admin/HR logged in", "1. Go to /employees/add\n2. Fill all valid fields\n3. Submit form", "Name: Kevin Durant, Email: kevin@workforce.test, Phone: 9876543299, Salary: 90000", "Success message 'Employee created successfully' displayed", "Employee created successfully", "Pass", "High", "P1"),
        ("TC_VAL_026", "Validation", "Add Employee - Missing Required Field", "On /employees/add form", "1. Leave First Name empty\n2. Fill other fields\n3. Submit", "First Name: Empty", "Validation error 'First Name is required' displayed", "Displayed 'First Name is required'", "Pass", "High", "P1"),
        ("TC_VAL_027", "Validation", "Add Employee - Invalid Email Format", "On /employees/add form", "1. Enter 'invalid-email-string'\n2. Submit form", "Email: invalid-email-string", "Validation error 'Invalid email format' displayed", "Displayed 'Invalid email format'", "Pass", "High", "P1"),
        ("TC_VAL_028", "Validation", "Add Employee - Invalid Phone Format", "On /employees/add form", "1. Enter '123' in phone input\n2. Submit form", "Phone: 123", "Validation error 'Invalid phone number format' displayed", "Displayed 'Invalid phone number format'", "Pass", "High", "P1"),
        ("TC_VAL_029", "Validation", "Add Employee - Non-Numeric Salary", "On /employees/add form", "1. Enter 'ABC' in salary field\n2. Submit form", "Salary: ABC", "Validation error 'Salary must be a valid positive number'", "Displayed 'Salary must be a valid positive number'", "Pass", "High", "P1"),
        ("TC_VAL_030", "Validation", "Add Employee - Invalid Joining Date Format", "On /employees/add form", "1. Enter '31-02-2026' in joining date\n2. Submit form", "Date: 31-02-2026", "Validation error 'Invalid joining date format' displayed", "Displayed 'Invalid joining date format'", "Pass", "Medium", "P2"),
        ("TC_VAL_031", "Validation", "Add Employee - Duplicate Email Check", "On /employees/add form", "1. Enter existing email 'admin@workforce.test'\n2. Submit form", "Email: admin@workforce.test", "Validation error 'Employee with this email already exists'", "Displayed 'Employee with this email already exists'", "Pass", "High", "P1"),
        ("TC_CRUD_032", "Add/Edit", "Edit Employee Information", "Admin/HR logged in", "1. Click Edit on EMP003\n2. Change position to 'Senior QA Specialist'\n3. Save", "Position: Senior QA Specialist", "Redirect to details page with success banner and updated info", "Updated successfully", "Pass", "High", "P1"),
        ("TC_VAL_033", "Validation", "Edit Employee - Duplicate Email Conflict", "On edit employee page", "1. Change email to 'hr@workforce.test'\n2. Submit form", "Email: hr@workforce.test", "Validation error 'Employee with this email already exists'", "Displayed 'Employee with this email already exists'", "Pass", "High", "P1"),
        ("TC_CRUD_034", "Delete", "Delete Employee - Confirm Delete Path", "Admin logged in", "1. Click Delete on test employee row\n2. Modal appears\n3. Click 'Delete'", "Employee ID: EMP006", "Employee record deleted from DB and table", "Record deleted successfully", "Pass", "High", "P1"),
        ("TC_CRUD_035", "Delete", "Delete Employee - Cancel Modal Path", "Admin logged in", "1. Click Delete on employee row\n2. Modal appears\n3. Click 'Cancel'", "Employee ID: EMP001", "Modal closes, employee record remains intact", "Modal closed, record retained", "Pass", "Medium", "P2"),

        # ROLE-BASED ACCESS CONTROL (RBAC)
        ("TC_RBAC_036", "Role Access", "Admin Role Access Privileges", "Admin logged in", "1. Verify navbar links\n2. Verify delete buttons visible", "Role: Admin", "Full access to View, Add, Edit, Delete, Upload Docs", "Verified full access", "Pass", "High", "P1"),
        ("TC_RBAC_037", "Role Access", "HR Manager Delete Button Absence", "HR Manager logged in", "1. Navigate to /employees\n2. Inspect Actions column", "Role: HR Manager", "Add and Edit buttons visible; Delete buttons HIDDEN", "Delete buttons hidden", "Pass", "High", "P1"),
        ("TC_RBAC_038", "Role Access", "HR Manager Direct Delete API Attempt", "HR Manager logged in", "1. Trigger POST /employees/1/delete directly", "POST /employees/1/delete", "Request denied with 302/403 'Only Admin can delete employees'", "Access forbidden error banner displayed", "Pass", "High", "P1"),
        ("TC_RBAC_039", "Role Access", "Employee Role Navigation Restriction", "Employee logged in", "1. Attempt navigating to /employees or /dashboard", "URL: /employees", "Redirected to /profile automatically", "Redirected to /profile", "Pass", "High", "P1"),
        ("TC_RBAC_040", "Role Access", "Employee Role Direct Access to Other Profile", "Employee logged in", "1. Attempt navigating to /employees/1 (Alexander Wright)", "URL: /employees/1", "Access denied, redirected to /profile", "Redirected to /profile", "Pass", "High", "P1"),

        # FILE UPLOAD MODULE
        ("TC_DOC_041", "Document", "Upload Valid PDF Document", "Admin/HR on employee detail page", "1. Select PDF file\n2. Click Upload button", "File: Resume.pdf", "File uploaded, listed in Uploaded Documents section", "PDF uploaded successfully", "Pass", "High", "P1"),
        ("TC_DOC_042", "Document", "Upload Valid PNG Image Document", "Admin/HR on employee detail page", "1. Select PNG file\n2. Click Upload button", "File: ID_Proof.png", "File uploaded, listed in Uploaded Documents section", "PNG uploaded successfully", "Pass", "High", "P1"),
        ("TC_DOC_043", "Document", "Upload Valid JPG Image Document", "Admin/HR on employee detail page", "1. Select JPG file\n2. Click Upload button", "File: Photo.jpg", "File uploaded, listed in Uploaded Documents section", "JPG uploaded successfully", "Pass", "High", "P1"),
        ("TC_DOC_044", "Document", "Upload Disallowed Executable File", "Admin/HR on employee detail page", "1. Select script.exe file\n2. Click Upload button", "File: script.exe", "Error 'Invalid file format. Allowed formats: PDF, PNG, JPG' displayed", "Displayed invalid file format error", "Pass", "High", "P1"),
        ("TC_DOC_045", "Document", "View / Download Uploaded Document", "User logged in", "1. Navigate to /documents hub\n2. Click 'View File' link", "Doc: sample_resume_EMP003.pdf", "Document opens in new browser tab", "Document opened successfully", "Pass", "Medium", "P2")
    ]

    # Styling definitions
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="15803D")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Write Headers
    ws.append(headers)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Write Data
    for row_idx, tc in enumerate(test_cases, 2):
        for col_idx, val in enumerate(tc, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Highlight Status
            if col_idx == 9 and val == "Pass":
                cell.fill = pass_fill
                cell.font = pass_font
                cell.alignment = Alignment(horizontal="center", vertical="top")

    # Set row height
    ws.row_dimensions[1].height = 28
    for row in range(2, len(test_cases) + 2):
        ws.row_dimensions[row].height = 45

    # Set column widths
    col_widths = {
        1: 14, # ID
        2: 15, # Module
        3: 35, # Scenario
        4: 25, # Preconditions
        5: 35, # Steps
        6: 30, # Data
        7: 35, # Expected
        8: 30, # Actual
        9: 12, # Status
        10: 12, # Severity
        11: 10  # Priority
    }

    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    os.makedirs("manual_testing", exist_ok=True)
    file_path = os.path.join("manual_testing", "test_cases.xlsx")
    wb.save(file_path)
    print(f"Excel file generated successfully with {len(test_cases)} test cases at: {file_path}")

if __name__ == "__main__":
    generate_manual_test_cases()
